from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl.drawing.image import Image as XLImage

from workbook_utils import autosize_columns, ensure_dir, new_workbook, utc_now_iso, write_kv_sheet, write_table


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
OUTPUTS_DIR = ROOT / "outputs"
THUMBNAIL_DIR = ROOT / ".scratch" / "card-thumbnails"
REPORT_PATH = DATA_DIR / "gundam_card_keep_report.json"
OUTPUT_PATH = OUTPUTS_DIR / "gundam-card-keep-priority.xlsx"


COLOR_SHEETS = ["Blue", "Green", "Purple", "Red", "White"]


def safe_float(value: Any, default: float = 0.0) -> float:
    try:
        if value in (None, ""):
            return default
        return float(str(value))
    except (TypeError, ValueError):
        return default


def thumbnail_path(card_number: str) -> Path | None:
    if not card_number:
        return None
    candidate = THUMBNAIL_DIR / f"{card_number}.jpg"
    if candidate.exists():
        return candidate
    return None


def sort_keep_rows(row: dict[str, Any]) -> tuple[float, float, str]:
    return (
        safe_float(row.get("keep_score"), 0.0),
        safe_float(row.get("raw_keep_signal"), 0.0),
        str(row.get("card_number") or ""),
    )


def write_keep_sheet(
    ws,
    rows: list[dict[str, Any]],
    columns: list[str],
    *,
    include_thumbnails: bool = True,
) -> dict[str, int]:
    missing_thumbnails = 0
    embedded_thumbnails = 0

    if include_thumbnails:
        ws.append(["thumbnail", *columns])
    else:
        ws.append(columns)

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 18
    if include_thumbnails:
        ws.column_dimensions["A"].width = 14

    for idx, row in enumerate(rows, start=2):
        if include_thumbnails:
            ws.row_dimensions[idx].height = 62
            ws.append(["", *[row.get(col, "") for col in columns]])
            card_number = str(row.get("card_number") or "")
            path = thumbnail_path(card_number)
            if path is None:
                missing_thumbnails += 1
                continue
            try:
                img = XLImage(str(path))
                img.width = 80
                img.height = 56
                ws.add_image(img, f"A{idx}")
                embedded_thumbnails += 1
            except Exception:  # noqa: BLE001
                missing_thumbnails += 1
        else:
            ws.append([row.get(col, "") for col in columns])

    autosize_columns(ws)
    return {"embedded_thumbnails": embedded_thumbnails, "missing_thumbnails": missing_thumbnails}


def main() -> int:
    if not REPORT_PATH.exists():
        raise SystemExit(f"Missing report: {REPORT_PATH}")

    report: dict[str, Any] = json.loads(REPORT_PATH.read_text(encoding="utf-8"))
    generated_at_utc = report.get("generated_at_utc") or utc_now_iso()
    latest_event_date = report.get("latest_event_date") or ""

    all_cards: list[dict[str, Any]] = report.get("all_cards") or []
    top_by_color: list[dict[str, Any]] = report.get("top_by_color") or []
    low_signal: list[dict[str, Any]] = report.get("low_signal_main_deck") or []
    color_summary: list[dict[str, Any]] = report.get("color_summary") or []

    wb = new_workbook()
    summary = wb.create_sheet("Summary")
    write_kv_sheet(
        summary,
        [
            ("generated_at_utc", generated_at_utc),
            ("latest_event_date", latest_event_date),
            ("all_cards", len(all_cards)),
            ("top_by_color", len(top_by_color)),
            ("low_signal_main_deck", len(low_signal)),
            ("thumbnail_dir", str(THUMBNAIL_DIR)),
        ],
    )

    ws = wb.create_sheet("Top By Color")
    thumb_stats_top = write_keep_sheet(
        ws,
        top_by_color,
        [
            "color",
            "card_number",
            "card_name",
            "keep_tier",
            "keep_score",
            "deck_count",
            "event_count",
            "archetype_count",
            "last_seen",
            "rarity",
            "package_codes",
            "market_price_usd",
            "market_price_brl_x5",
            "price_status",
            "price_error",
            "price_last_checked_utc",
        ],
        include_thumbnails=True,
    )

    for color in COLOR_SHEETS:
        rows = [r for r in all_cards if str(r.get("color") or "") == color]
        rows.sort(key=sort_keep_rows, reverse=True)
        ws = wb.create_sheet(color)
        write_keep_sheet(
            ws,
            rows,
            [
                "card_number",
                "card_name",
                "keep_tier",
                "keep_score",
                "keep_recommendation",
                "deck_count",
                "event_count",
                "archetype_count",
                "last_seen",
                "rarity",
                "package_codes",
                "package_labels",
                "market_price_usd",
                "market_price_brl_x5",
                "price_status",
                "price_error",
                "price_last_checked_utc",
            ],
            include_thumbnails=True,
        )

    other_rows = [r for r in all_cards if str(r.get("color") or "") not in COLOR_SHEETS]
    other_rows.sort(key=sort_keep_rows, reverse=True)
    ws = wb.create_sheet("Other Colors")
    write_keep_sheet(
        ws,
        other_rows,
        [
            "color",
            "card_number",
            "card_name",
            "keep_tier",
            "keep_score",
            "keep_recommendation",
            "deck_count",
            "event_count",
            "archetype_count",
            "last_seen",
            "rarity",
            "package_codes",
            "market_price_usd",
            "market_price_brl_x5",
            "price_status",
            "price_error",
            "price_last_checked_utc",
        ],
        include_thumbnails=True,
    )

    ws = wb.create_sheet("All Rankings")
    all_sorted = list(all_cards)
    all_sorted.sort(key=sort_keep_rows, reverse=True)
    thumb_stats_all = write_keep_sheet(
        ws,
        all_sorted,
        [
            "color",
            "card_number",
            "card_name",
            "keep_tier",
            "keep_score",
            "keep_recommendation",
            "deck_count",
            "event_count",
            "archetype_count",
            "last_seen",
            "rarity",
            "package_codes",
            "market_price_usd",
            "market_price_brl_x5",
            "price_status",
            "price_error",
            "price_last_checked_utc",
        ],
        include_thumbnails=True,
    )

    ws = wb.create_sheet("Low Signal")
    low_sorted = list(low_signal)
    low_sorted.sort(key=sort_keep_rows, reverse=True)
    write_keep_sheet(
        ws,
        low_sorted,
        [
            "color",
            "card_number",
            "card_name",
            "keep_tier",
            "keep_score",
            "keep_recommendation",
            "last_seen",
            "rarity",
            "market_price_usd",
            "market_price_brl_x5",
            "price_status",
        ],
        include_thumbnails=True,
    )

    ws = wb.create_sheet("Color Summary")
    write_table(ws, color_summary, None)

    ws = wb.create_sheet("Scoring Model")
    notes = report.get("scoring_notes") or []
    tiers = report.get("tier_definitions") or []
    ws.append(["generated_at_utc", generated_at_utc])
    ws.append(["latest_event_date", latest_event_date])
    ws.append([])
    ws.append(["tier", "meaning", "default_action"])
    for row in tiers:
        ws.append([row.get("tier", ""), row.get("meaning", ""), row.get("default_action", "")])
    ws.append([])
    ws.append(["notes"])
    for note in notes:
        ws.append([str(note)])
    autosize_columns(ws)

    ensure_dir(OUTPUTS_DIR)
    wb.save(OUTPUT_PATH)

    embedded = thumb_stats_top["embedded_thumbnails"] + thumb_stats_all["embedded_thumbnails"]
    missing = thumb_stats_top["missing_thumbnails"] + thumb_stats_all["missing_thumbnails"]
    print(
        json.dumps(
            {
                "workbook": str(OUTPUT_PATH),
                "generated_at_utc": generated_at_utc,
                "latest_event_date": latest_event_date,
                "thumbnail_embed_summary": {
                    "embedded_thumbnails_sampled": embedded,
                    "missing_thumbnails_sampled": missing,
                },
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

