from __future__ import annotations

from collections.abc import Iterable
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def utc_now_iso() -> str:
    return datetime.utcnow().replace(microsecond=0).isoformat() + "Z"


def autosize_columns(ws, max_width: int = 60, min_width: int = 10) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows(values_only=True):
        for idx, value in enumerate(row, start=1):
            if value is None:
                continue
            value_str = str(value)
            widths[idx] = max(widths.get(idx, 0), len(value_str))
    for idx, length in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(min_width, min(max_width, length + 2))


def write_table(
    ws,
    rows: list[dict[str, Any]],
    columns: list[str] | None = None,
    *,
    header_fill: str = "10243E",
    header_font_color: str = "FFFFFF",
) -> None:
    if columns is None:
        keys: set[str] = set()
        for row in rows:
            keys.update(row.keys())
        columns = sorted(keys)

    ws.append(columns)
    header = ws[1]
    for cell in header:
        cell.font = Font(bold=True, color=header_font_color)
        cell.fill = PatternFill("solid", fgColor=header_fill)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        ws.append([row.get(col, "") for col in columns])

    ws.freeze_panes = "A2"
    autosize_columns(ws)


def write_kv_sheet(ws, items: Iterable[tuple[str, Any]]) -> None:
    ws.append(["key", "value"])
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)
    for key, value in items:
        ws.append([key, value if value is not None else ""])
    ws.freeze_panes = "A2"
    autosize_columns(ws)


def new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb

